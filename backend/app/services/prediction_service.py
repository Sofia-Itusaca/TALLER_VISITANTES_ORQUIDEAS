from pathlib import Path
from openpyxl import load_workbook

DATA_FILE = Path("/app/data/base_datos_clima_visitantes_orquideas.xlsx")


class PredictionService:
    """
    Modelo inicial basado en parámetros de ejemplo del archivo Excel.
    Los valores deben recalibrarse cuando existan datos reales.
    """

    def __init__(self):
        self.model = self._load_model()

    def _load_model(self):
        path = DATA_FILE
        if not path.exists():
            # Para ejecución local desde backend/
            path = Path(__file__).resolve().parents[3] / "data" / "base_datos_clima_visitantes_orquideas.xlsx"

        wb = load_workbook(path, data_only=True)
        ws = wb["Modelo"]

        model = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            clima, base, temp_factor, rain_factor, weekend_factor = row[:5]

            try:
                model[str(clima)] = {
                    "base": float(base),
                    "temp_factor": float(temp_factor),
                    "rain_factor": float(rain_factor),
                    "weekend_factor": float(weekend_factor),
                }
            except (TypeError, ValueError):
                continue
        return model

    def normalize_condition(self, description: str) -> str:
        text = description.lower()

        if "torment" in text or "lluvia" in text or "rain" in text:
            # Lluvia ligera si el pronóstico tiene poca precipitación; se ajusta después.
            return "Lluvia"

        if "soleado" in text or "despejado" in text or "sunny" in text or "clear" in text:
            return "Soleado"

        if "parcial" in text or "parcialmente" in text:
            return "Parcialmente nublado"

        if "nublado" in text or "overcast" in text:
            return "Nublado"

        return "Nublado"

    def estimate(self, description: str, rain_mm: float, date_str: str) -> dict:
        import datetime

        condition = self.normalize_condition(description)

        # Diferenciamos lluvia ligera de lluvia fuerte.
        if condition == "Lluvia":
            condition = "Lluvia ligera" if rain_mm <= 10 else "Lluvia"

        params = self.model.get(condition, self.model["Nublado"])
        d = datetime.date.fromisoformat(date_str)
        weekend = d.weekday() >= 5

        estimate = (
            params["base"]
            * params["temp_factor"]
            * params["rain_factor"]
            * (params["weekend_factor"] if weekend else 1.0)
        )

        estimate = max(5, round(estimate))
        low = max(0, round(estimate * 0.85))
        high = round(estimate * 1.15)

        return {
            "conditionModel": condition,
            "estimatedVisitors": estimate,
            "rangeMin": low,
            "rangeMax": high,
        }
